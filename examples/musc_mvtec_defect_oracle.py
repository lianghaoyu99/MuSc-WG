"""Run a score-calibrated defect-type Mani oracle on MVTec.

The command-line interface intentionally mirrors ``examples/musc_main.py``::

    python .\examples\musc_mvtec_defect_oracle.py

Each product is first evaluated with lambda=0 to establish one shared score
scale. Predictions from every nonzero defect-specific lambda/tau pair are then
quantile-mapped using only the common ``good`` images before the selected defect
images are assembled. Normal images and defects configured with lambda=0 use
the baseline prediction directly. The script recomputes metrics over the
complete product category, macro-averages the 15 products exactly like
``MuSc.main()``, and writes ``results.xlsx`` plus the final selected ``vis``.

Important: this is an oracle experiment. The true defect type and parameters
selected on the test set are used, so the result is an analysis upper bound,
not a deployable or fair main benchmark result.

Completed product categories are saved to CSV. Rerun the same command after an
interruption to resume from the next unfinished product.
"""

import argparse
import csv
import gc
import hashlib
import json
import os
import sys
from collections import defaultdict
from copy import deepcopy
from pathlib import Path

import numpy as np
import torch
from openpyxl import Workbook

sys.path.append(os.getcwd())

from models.musc import MuSc
from utils.load_config import load_yaml
from utils.metrics import compute_metrics


METRIC_NAMES = [
    "image_auroc",
    "image_f1",
    "image_ap",
    "pixel_auroc",
    "pixel_f1",
    "pixel_ap",
    "aupro",
]

RESULT_FIELDS = [
    "signature",
    "category",
    *METRIC_NAMES,
    "unique_parameter_pairs",
    "sum_ms_per_image",
    "max_gpu_allocated_mb",
    "max_gpu_reserved_mb",
    "calibration_method",
    "calibration_quantiles",
    "calibration_pixel_samples",
]


def parse_args():
    parser = argparse.ArgumentParser(
        description="MVTec defect-level Mani oracle evaluation"
    )
    parser.add_argument(
        "--config",
        default="./configs/musc_mvtec_defect_oracle.yaml",
        help="Configuration containing manifold_parameters_by_defect.",
    )
    parser.add_argument("--data_path", default=None)
    parser.add_argument("--device", type=int, default=None)
    parser.add_argument("--output_dir", default=None)
    parser.add_argument("--vis", choices=["true", "false"], default=None)
    parser.add_argument(
        "--save_excel", choices=["true", "false"], default=None
    )
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument(
        "--no_resume",
        action="store_true",
        help="Recompute all products instead of resuming the saved CSV.",
    )
    parser.add_argument(
        "--validate_only",
        action="store_true",
        help="Validate data folders and all defect parameters, then exit.",
    )
    return parser.parse_args()


def read_csv(path):
    if not path.is_file():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv_atomic(path, rows):
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=RESULT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    os.replace(temporary, path)


def make_signature(cfg, parameters, data_path, seed):
    value = {
        "protocol": (
            "baseline_good_and_lambda0_defects;"
            "good_reference_quantile_calibrated_nonzero_defects"
        ),
        "parameters": parameters,
        "data_path": str(Path(data_path).resolve()),
        "seed": seed,
        "datasets": {
            key: cfg["datasets"].get(key)
            for key in ("img_resize", "divide_num")
        },
        "models": cfg["models"],
        "testing": {
            key: cfg["testing"].get(key)
            for key in (
                "use_rscin",
                "manifold_level",
                "manifold_fusion",
                "manifold_normalize",
                "manifold_layers",
                "manifold_r_list",
                "manifold_grid_size",
                "manifold_ref_chunk_size",
                "defect_oracle_calibration",
            )
        },
    }
    encoded = json.dumps(value, sort_keys=True).encode("utf-8")
    return hashlib.sha1(encoded).hexdigest()[:12]


def resolve_pair(value, category, defect):
    if not isinstance(value, dict):
        raise ValueError(
            f"{category}/{defect} must map to temperature and lambda_m."
        )
    temperature = float(value.get("temperature", value.get("tau", 0.5)))
    lambda_m = float(value["lambda_m"])
    if temperature <= 0 or lambda_m < 0:
        raise ValueError(
            f"Invalid pair for {category}/{defect}: "
            f"temperature={temperature}, lambda_m={lambda_m}."
        )
    return lambda_m, temperature


def defect_names(paths):
    return np.asarray(
        [os.path.basename(os.path.dirname(path)) for path in paths]
    )


def sampled_flat_values(values, max_samples):
    """Return a deterministic, bounded sample without flattening a copy."""
    flattened = np.asarray(values).reshape(-1)
    if max_samples <= 0 or flattened.size <= max_samples:
        return flattened
    indices = np.linspace(
        0, flattened.size - 1, num=max_samples, dtype=np.int64
    )
    return flattened[indices]


def quantile_anchors(source, target, quantile_count, max_samples):
    """Build a monotonic map from one score distribution to another."""
    source_values = sampled_flat_values(source, max_samples)
    target_values = sampled_flat_values(target, max_samples)
    levels = np.linspace(0.0, 1.0, num=quantile_count)
    source_quantiles = np.quantile(source_values, levels).astype(np.float64)
    target_quantiles = np.quantile(target_values, levels).astype(np.float64)

    # Ensure the tails cover the complete arrays rather than only the sample.
    source_quantiles[0] = float(np.min(source))
    source_quantiles[-1] = float(np.max(source))
    target_quantiles[0] = float(np.min(target))
    target_quantiles[-1] = float(np.max(target))

    unique_source, inverse = np.unique(
        source_quantiles, return_inverse=True
    )
    target_sum = np.zeros(unique_source.shape, dtype=np.float64)
    target_count = np.zeros(unique_source.shape, dtype=np.int64)
    np.add.at(target_sum, inverse, target_quantiles)
    np.add.at(target_count, inverse, 1)
    unique_target = target_sum / target_count
    return unique_source, unique_target


def quantile_calibrate(
    selected_values,
    source_distribution,
    target_distribution,
    quantile_count,
    max_samples,
):
    """Map selected scores onto a shared baseline distribution."""
    source_anchors, target_anchors = quantile_anchors(
        source_distribution,
        target_distribution,
        quantile_count,
        max_samples,
    )
    if source_anchors.size == 1:
        return np.full_like(
            selected_values,
            target_anchors[0],
            dtype=np.float64,
        )
    calibrated = np.interp(
        np.asarray(selected_values, dtype=np.float64),
        source_anchors,
        target_anchors,
    )
    # np.interp clamps tails. Linear extrapolation keeps scores above the
    # normal range distinguishable instead of collapsing anomalies to the
    # maximum score observed in good images.
    selected_array = np.asarray(selected_values, dtype=np.float64)
    below = selected_array < source_anchors[0]
    above = selected_array > source_anchors[-1]
    if np.any(below):
        lower_slope = (
            (target_anchors[1] - target_anchors[0])
            / (source_anchors[1] - source_anchors[0])
        )
        calibrated[below] = target_anchors[0] + lower_slope * (
            selected_array[below] - source_anchors[0]
        )
    if np.any(above):
        upper_slope = (
            (target_anchors[-1] - target_anchors[-2])
            / (source_anchors[-1] - source_anchors[-2])
        )
        calibrated[above] = target_anchors[-1] + upper_slope * (
            selected_array[above] - source_anchors[-1]
        )
    return calibrated


def discover_defects(data_path, category):
    test_dir = Path(data_path) / category / "test"
    if not test_dir.is_dir():
        raise ValueError(f"MVTec test directory does not exist: {test_dir}")
    return {
        path.name for path in test_dir.iterdir() if path.is_dir()
    }


def metric_mean(rows):
    return {
        metric: float(np.mean([float(row[metric]) for row in rows]))
        for metric in METRIC_NAMES
    }


def print_metrics(category, row):
    print(category)
    print(
        "image-level, auroc:{}, f1:{}, ap:{}".format(
            float(row["image_auroc"]) * 100,
            float(row["image_f1"]) * 100,
            float(row["image_ap"]) * 100,
        )
    )
    print(
        "pixel-level, auroc:{}, f1:{}, ap:{}, aupro:{}".format(
            float(row["pixel_auroc"]) * 100,
            float(row["pixel_f1"]) * 100,
            float(row["pixel_ap"]) * 100,
            float(row["aupro"]) * 100,
        )
    )


def write_results_xlsx(path, categories, result_map, include_mean):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MuSc_results"
    headers = [
        "auroc_px",
        "f1_px",
        "ap_px",
        "aupro",
        "auroc_sp",
        "f1_sp",
        "ap_sp",
    ]
    for column, header in enumerate(headers, start=2):
        sheet.cell(row=1, column=column, value=header)

    completed_rows = []
    output_row = 2
    for category in categories:
        if category not in result_map:
            continue
        row = result_map[category]
        completed_rows.append(row)
        sheet.cell(row=output_row, column=1, value=category)
        values = [
            row["pixel_auroc"],
            row["pixel_f1"],
            row["pixel_ap"],
            row["aupro"],
            row["image_auroc"],
            row["image_f1"],
            row["image_ap"],
        ]
        for column, value in enumerate(values, start=2):
            sheet.cell(
                row=output_row, column=column, value=float(value) * 100
            )
        output_row += 1

    if include_mean and completed_rows:
        means = metric_mean(completed_rows)
        sheet.cell(row=output_row, column=1, value="mean")
        values = [
            means["pixel_auroc"],
            means["pixel_f1"],
            means["pixel_ap"],
            means["aupro"],
            means["image_auroc"],
            means["image_f1"],
            means["image_ap"],
        ]
        for column, value in enumerate(values, start=2):
            sheet.cell(
                row=output_row, column=column, value=float(value) * 100
            )
    workbook.save(path)


def main():
    args = parse_args()
    cfg = deepcopy(load_yaml(args.config))
    cfg["datasets"]["dataset_name"] = "mvtec_ad"
    cfg["datasets"]["class_name"] = "ALL"
    if args.data_path is not None:
        cfg["datasets"]["data_path"] = args.data_path
    if args.device is not None:
        cfg["device"] = str(args.device)
    if args.output_dir is not None:
        cfg["testing"]["output_dir"] = args.output_dir
    if args.vis is not None:
        cfg["testing"]["vis"] = args.vis == "true"
    if args.save_excel is not None:
        cfg["testing"]["save_excel"] = args.save_excel == "true"

    data_path = cfg["datasets"]["data_path"]
    if not Path(data_path).is_dir():
        raise ValueError(f"MVTec data path does not exist: {data_path}")

    all_parameters = cfg["testing"].get(
        "manifold_parameters_by_defect", {}
    ).get("mvtec_ad", {})
    if not all_parameters:
        raise ValueError(
            "testing.manifold_parameters_by_defect.mvtec_ad is empty."
        )
    categories = list(all_parameters)
    calibration = cfg["testing"].get("defect_oracle_calibration", {}) or {}
    calibration_method = calibration.get("method", "quantile")
    calibration_quantiles = int(calibration.get("quantiles", 257))
    calibration_pixel_samples = int(
        calibration.get("pixel_samples", 1000000)
    )
    baseline_temperature = float(
        calibration.get("baseline_temperature", 0.5)
    )
    if calibration_method != "quantile":
        raise ValueError(
            "Only defect_oracle_calibration.method=quantile is supported."
        )
    if calibration_quantiles < 3:
        raise ValueError("Calibration quantiles must be at least 3.")
    if calibration_pixel_samples < calibration_quantiles:
        raise ValueError(
            "Calibration pixel_samples must be at least quantiles."
        )
    if baseline_temperature <= 0:
        raise ValueError("Calibration baseline_temperature must be positive.")

    # Validate the fine-grained mapping before loading the backbone.
    for category, configured in all_parameters.items():
        available = discover_defects(data_path, category)
        if "good" not in available:
            raise ValueError(f"{category} does not contain test/good.")
        if "__all__" not in configured:
            raise ValueError(f"{category} is missing the __all__ pair for good.")
        configured_defects = set(configured) - {"__all__"}
        disk_defects = available - {"good"}
        if configured_defects != disk_defects:
            raise ValueError(
                f"{category} defect mapping mismatch; "
                f"missing={sorted(disk_defects - configured_defects)}, "
                f"extra={sorted(configured_defects - disk_defects)}."
            )
        for defect, value in configured.items():
            resolve_pair(value, category, defect)

    if args.validate_only:
        defect_count = sum(len(values) - 1 for values in all_parameters.values())
        print(
            f"Validation passed: {len(categories)} products, "
            f"{defect_count} defect types, and one __all__ pair per product. "
            f"Calibration={calibration_method}, "
            f"quantiles={calibration_quantiles}, "
            f"pixel_samples={calibration_pixel_samples}."
        )
        return

    save_vis = bool(cfg["testing"].get("vis", True))
    save_excel = bool(cfg["testing"].get("save_excel", True))
    if cfg["testing"].get("vis_tsne", False):
        print(
            "vis_tsne is disabled for this oracle script because features "
            "cannot be stitched across parameter runs.",
            flush=True,
        )
    # Intermediate parameter runs must not write misleading visualizations.
    cfg["testing"]["vis"] = False
    cfg["testing"]["vis_tsne"] = False
    cfg["testing"]["save_excel"] = False
    cfg["testing"]["use_manifold"] = True

    artifact_dir = (
        Path(cfg["testing"]["output_dir"])
        / "mvtec_ad"
        / cfg["models"]["backbone_name"]
        / f"imagesize{cfg['datasets']['img_resize']}"
    )
    artifact_dir.mkdir(parents=True, exist_ok=True)
    results_csv = artifact_dir / "defect_oracle_category_results.csv"
    results_xlsx = artifact_dir / "results.xlsx"
    signature = make_signature(
        cfg, all_parameters, data_path, args.seed
    )

    saved = [] if args.no_resume else read_csv(results_csv)
    result_map = {
        row["category"]: row
        for row in saved
        if row.get("signature") == signature
        and row.get("category") in categories
    }
    pending = [category for category in categories if category not in result_map]
    print(
        "WARNING: score-calibrated defect-label oracle; not a deployable "
        "benchmark setting.",
        flush=True,
    )
    print(
        f"Signature: {signature}; completed={len(result_map)}, "
        f"pending={len(pending)}",
        flush=True,
    )

    model = None
    if pending:
        model = MuSc(cfg, seed=args.seed)
        dataset_values = model.manifold_parameters_by_dataset.setdefault(
            "mvtec_ad", {}
        )
        category_values = dataset_values.setdefault("categories", {})

    try:
        for category_index, category in enumerate(categories, start=1):
            if category in result_map:
                print(
                    f"[{category_index}/{len(categories)}] "
                    f"Resume skip: {category}",
                    flush=True,
                )
                continue

            assignments = {
                defect: resolve_pair(value, category, defect)
                for defect, value in all_parameters[category].items()
                if defect != "__all__"
            }
            grouped = defaultdict(list)
            baseline_defects = []
            for defect, pair in assignments.items():
                if pair[0] == 0.0:
                    baseline_defects.append(defect)
                else:
                    grouped[pair].append(defect)
            total_parameter_pairs = 1 + len(grouped)

            print(
                f"\n=== [{category_index}/{len(categories)}] {category}: "
                f"{total_parameter_pairs} runs including shared baseline ===",
                flush=True,
            )
            sum_ms = 0.0
            max_allocated = 0.0
            max_reserved = 0.0

            # Establish one common scale for the complete product category.
            print(
                f"pair [1/{total_parameter_pairs}]: lambda=0.0, "
                f"tau={baseline_temperature}, use=good"
                + (
                    "," + ",".join(sorted(baseline_defects))
                    if baseline_defects
                    else ""
                ),
                flush=True,
            )
            category_values[category] = {
                "lambda_m": 0.0,
                "temperature": baseline_temperature,
            }
            (
                _,
                _,
                ms_per_image,
                gpu_allocated_mb,
                gpu_reserved_mb,
                baseline_prediction,
            ) = model.make_category_data(
                category,
                return_predictions=True,
                compute_metrics_output=False,
            )
            paths = list(baseline_prediction["image_paths"])
            current_defects = defect_names(paths)
            labels = np.asarray(
                baseline_prediction["image_labels"], dtype=np.int8
            ).copy()
            masks = np.asarray(
                baseline_prediction["pixel_masks"], dtype=np.uint8
            ).copy()
            baseline_image_scores = np.asarray(
                baseline_prediction["image_scores"], dtype=np.float64
            ).copy()
            baseline_pixel_scores = np.asarray(
                baseline_prediction["pixel_scores"], dtype=np.float32
            ).copy()
            final_image_scores = baseline_image_scores.copy()
            final_pixel_scores = baseline_pixel_scores.copy()
            good = current_defects == "good"
            filled = good.copy()
            if baseline_defects:
                filled |= np.isin(current_defects, baseline_defects)
            sum_ms += float(ms_per_image)
            max_allocated = max(max_allocated, float(gpu_allocated_mb))
            max_reserved = max(max_reserved, float(gpu_reserved_mb))
            del baseline_prediction
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()

            for pair_index, ((lambda_m, temperature), selected_defects) in enumerate(
                sorted(grouped.items()), start=2
            ):
                print(
                    f"pair [{pair_index}/{total_parameter_pairs}]: "
                    f"lambda={lambda_m}, tau={temperature}, "
                    f"use={','.join(sorted(selected_defects))}",
                    flush=True,
                )
                category_values[category] = {
                    "lambda_m": lambda_m,
                    "temperature": temperature,
                }
                (
                    _,
                    _,
                    ms_per_image,
                    gpu_allocated_mb,
                    gpu_reserved_mb,
                    prediction,
                ) = model.make_category_data(
                    category,
                    return_predictions=True,
                    compute_metrics_output=False,
                )

                current_paths = list(prediction["image_paths"])
                current_defects = defect_names(current_paths)
                if current_paths != paths:
                    raise RuntimeError(
                        f"{category}: image order changed between runs."
                    )
                if not np.array_equal(
                    labels,
                    np.asarray(prediction["image_labels"], dtype=np.int8),
                ):
                    raise RuntimeError(
                        f"{category}: labels changed between runs."
                    )

                selected = np.isin(current_defects, selected_defects)
                if np.any(filled & selected):
                    raise RuntimeError(
                        f"{category}: some images were assigned twice."
                    )
                # Learn the shared scale only from common normal images. The
                # selected anomaly values never influence the calibration map.
                final_image_scores[selected] = quantile_calibrate(
                    prediction["image_scores"][selected],
                    prediction["image_scores"][good],
                    baseline_image_scores[good],
                    calibration_quantiles,
                    max_samples=0,
                )
                final_pixel_scores[selected] = quantile_calibrate(
                    prediction["pixel_scores"][selected],
                    prediction["pixel_scores"][good],
                    baseline_pixel_scores[good],
                    calibration_quantiles,
                    calibration_pixel_samples,
                )
                filled[selected] = True
                sum_ms += float(ms_per_image)
                max_allocated = max(max_allocated, float(gpu_allocated_mb))
                max_reserved = max(max_reserved, float(gpu_reserved_mb))

                del prediction
                gc.collect()
                if torch.cuda.is_available():
                    torch.cuda.empty_cache()

            if not np.all(filled):
                missing = sorted(set(defect_names(paths)[~filled]))
                raise RuntimeError(
                    f"{category}: no selected prediction for {missing}."
                )

            image_metric, pixel_metric = compute_metrics(
                labels, final_image_scores, masks, final_pixel_scores
            )
            row = {
                "signature": signature,
                "category": category,
                "image_auroc": float(image_metric[0]),
                "image_f1": float(image_metric[1]),
                "image_ap": float(image_metric[2]),
                "pixel_auroc": float(pixel_metric[0]),
                "pixel_f1": float(pixel_metric[1]),
                "pixel_ap": float(pixel_metric[2]),
                "aupro": float(pixel_metric[3]),
                "unique_parameter_pairs": total_parameter_pairs,
                # This is oracle construction cost, not deployment latency.
                "sum_ms_per_image": sum_ms,
                "max_gpu_allocated_mb": max_allocated,
                "max_gpu_reserved_mb": max_reserved,
                "calibration_method": calibration_method,
                "calibration_quantiles": calibration_quantiles,
                "calibration_pixel_samples": calibration_pixel_samples,
            }
            result_map[category] = row

            if save_vis:
                print("visualization...")
                model.visualization(
                    paths,
                    labels.tolist(),
                    final_pixel_scores,
                    masks,
                    category,
                )
            ordered = [
                result_map[name] for name in categories if name in result_map
            ]
            write_csv_atomic(results_csv, ordered)
            if save_excel:
                write_results_xlsx(
                    results_xlsx,
                    categories,
                    result_map,
                    include_mean=len(result_map) == len(categories),
                )
            print("DEFECT-ORACLE CATEGORY RESULT")
            print_metrics(category, row)

            del (
                masks,
                final_pixel_scores,
                final_image_scores,
                baseline_pixel_scores,
                baseline_image_scores,
            )
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()
    except KeyboardInterrupt:
        print(
            "\nInterrupted. Finished products are saved; rerun the same "
            "command to resume.",
            flush=True,
        )
        raise

    ordered = [result_map[category] for category in categories]
    means = metric_mean(ordered)
    if save_excel:
        write_results_xlsx(
            results_xlsx,
            categories,
            result_map,
            include_mean=True,
        )

    print("\n========== MuSc-style defect-oracle result ==========")
    for category in categories:
        print_metrics(category, result_map[category])
    print_metrics("mean", means)
    construction_time = float(
        np.mean([float(row["sum_ms_per_image"]) for row in ordered])
    )
    print(
        f"Oracle construction: {construction_time:.2f}ms per image "
        "(repeated parameter runs; not deployment latency)"
    )
    print(f"Category checkpoint: {results_csv}")
    if save_excel:
        print(f"Excel: {results_xlsx}")
    if save_vis:
        print(f"Visualization: {artifact_dir / 'vis'}")


if __name__ == "__main__":
    main()
